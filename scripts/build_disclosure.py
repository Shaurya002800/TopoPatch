"""Build the editable IDF-B continuation from measured local evidence."""
from pathlib import Path
import json
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image, ImageDraw, ImageFont
from pypdf import PdfReader
import textwrap
from io import BytesIO
import subprocess
from reportlab.pdfgen import canvas
from pypdf import PdfWriter
ROOT=Path(__file__).resolve().parents[1];OUT=ROOT/'deliverables';OUT.mkdir(exist_ok=True)
d=Document();sec=d.sections[0];sec.page_height=Inches(11.7);sec.page_width=Inches(8.3);sec.top_margin=sec.bottom_margin=Inches(.65);sec.left_margin=sec.right_margin=Inches(.7)
for name in ['Normal','Title','Heading 1','Heading 2','Caption']:
 s=d.styles[name];s.font.name='Arial';s.font.color.rgb=RGBColor(0,0,0)

for st in d.styles:
 for border in list(st.element.iter(qn('w:pBdr'))): border.getparent().remove(border)
d.styles['Normal'].font.name='Times New Roman';d.styles['Normal'].font.size=Pt(10.5);d.styles['Normal'].paragraph_format.space_after=Pt(7)
d.styles['Heading 1'].font.name='Times New Roman';d.styles['Heading 1'].font.size=Pt(12)
d.styles['Heading 2'].font.name='Times New Roman';d.styles['Heading 2'].font.size=Pt(11)
head=sec.header.paragraphs[0];head.text='©VIT IPR&TTCELL';head.style='Caption'
f=sec.footer.paragraphs[0];f.alignment=2;f.add_run('Page ');field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');f._p.append(field)
def p(t,bold=False):
 x=d.add_paragraph();x.add_run(t).bold=bold;return x
def h(t):d.add_heading(t,1)
def page():pass
def table(headers,rows):
 t=d.add_table(rows=1, cols=len(headers));t.style='Table Grid'
 for c,v in zip(t.rows[0].cells,headers):c.text=v
 for c in t.rows[0].cells:
  for r in c.paragraphs[0].runs:r.bold=True
  c.paragraphs[0].paragraph_format.keep_with_next=True
  sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'EDEDED');c._tc.get_or_add_tcPr().append(sh)
 for row in rows:
  for c,v in zip(t.add_row().cells,row):c.text=str(v)
 for row in t.rows:
  trPr=row._tr.get_or_add_trPr();cant=OxmlElement('w:cantSplit');trPr.append(cant)
  for c in row.cells:
   for pp in c.paragraphs:
    pp.paragraph_format.space_after=Pt(5)
    for r in pp.runs:r.font.size=Pt(9)
 return t

logo_path=ROOT/'artifacts_private/vit-idf-logo.png'
logo_path.write_bytes(PdfReader(ROOT/'doc/Invention_Disclosure_Format_B.pdf').pages[0].images[0].data)
mast=d.add_table(rows=1,cols=3);mast.style='Table Grid'
mast.cell(0,0).paragraphs[0].add_run().add_picture(str(logo_path),width=Inches(1.45))
mast.cell(0,1).text='Invention Disclosure Format (IDF)-B'
mast.cell(0,2).text='Document No.  02-IPR-R003\nIssue No/Date  2/01.02.2024\nAmd. No/Date  0/00.00.0000'
for cell in mast.rows[0].cells:
 for pp in cell.paragraphs:
  for run in pp.runs:run.font.name='Times New Roman';run.font.size=Pt(9);run.bold=True
p('')
h('1. Title of the invention:')
p('Validation and controlled commit of partial embedding model migrations',True)
h('2. Field /Area of invention:')
p('The invention concerns search systems that store document embeddings in a vector index. It focuses on replacing the embedding model for selected documents, checking the resulting search behaviour, and controlling when those changes become active. Possible uses include enterprise search and retrieval augmented generation systems.')
p('TopoPatch addresses a practical question: when only some documents have moved to a new embedding model, can the system check the effect on search results before making that change permanent? The prototype calculates the proposed mixed result, records the evidence and repeats the checks inside the commit operation.')
p('The present prototype shows this control sequence locally. It has not yet shown that selective migration gives the quality of a full upgrade at lower cost.')
p('The main embodiment removes migrated documents from the active old-model ranking. A separate experimental variant retains that ranking as a shadow signal. Their results and limitations are kept separate throughout this document.')
page()
h('3. Prior Patents and Publications from literature')
p('The references below show why model migration, selective indexing and validation cannot, by themselves, be treated as new. The final column identifies the narrower feature that needs comparison with each reference. These are candidate distinctions, not a finding that the invention is novel.')
table(['Reference','Relevant teaching','Proposed distinction to examine'],[
('[R1] US20240303516A1\nVector as a service\nHubSpot, 12 Sep 2024','Vector management and reindexing associated with embedding model changes.','Action-specific mixed-state retrieval evidence bound to source versions and checked again at commit.'),
('[R2] US20250342179A1\nCloudera, 6 Nov 2025','Detecting dataset mutations and selectively updating affected vector-index content.','Migration selection based on measured retrieval effects of an A-to-B representation change, including unchanged source content.'),
('[R3] Qdrant\nMigrate to a New Embedding Model\nOnline documentation','Operational model migration using new vectors while maintaining service.','A recorded partial-index action and its retrieval effects, checked again against source state at commit.'),
('[R8] US11610143B1\nAmazon, 21 Mar 2023','Client-specific tests guide whether an updated machine-learning model is deployed. Claims 1 to 3 cover validation-based migration decisions.','Test the exact partial A/B document-membership change and bind it to source state at commit. This narrower distinction still needs examination.'),
('[R9] US20260195572A1\n9 Jul 2026','Localized embedding patches with a model that maps new embeddings back to the original space; see claim 1.','TopoPatch uses native-space rankings without a trained cross-space mapping. Generic local patching is not claimed as new.'),
('[R10] SPFresh\nSOSP 2023','Incremental in-place vector-index updates using local partition rebalancing.','The prototype controls model migration through retrieval checks; it does not implement a new ANN repair method.'),
('[R12] Cormack et al.\nSIGIR 2009','Reciprocal rank fusion combines ranked result lists without cross-model score calibration.','Rank fusion is established supporting machinery, not the proposed novel step.'),
('[R11] SQLite documentation','Standard isolation and atomic database transactions.','The retrieval-specific admission logic needs examination; the database transaction itself is established.')])
p('BEIR and SciFact [R4] are evaluation resources used in Section 8, not competing inventions. Live web documentation needs a dated archived version before it is relied on as pre-filing prior art.')
p('The main question is whether testing a specific partial-index change and binding that result to a source-consistent commit adds a non-obvious technical contribution. The references make broad claims to validation, selective updates or rollback difficult to defend.',True)
page()
h('4. Summary and background of the invention (Address the gap / Novelty)')
p('Changing an embedding model can change what a search system returns. During a partial upgrade, some documents use the old model and others use the new one. Combining their result lists can push an important document down the ranking even when that document has not been edited. Testing only the new vectors can miss this problem.')
p('TopoPatch checks the proposed mixed index before activating the change. It uses the same ranking routine as the live search path and records which documents would move, their source hashes, the current migration epoch and the ranks of required results. The commit step checks that the evidence is still current and runs the retrieval checks again before updating active membership.')
p('The feature proposed for examination is the link between a tested partial-index change, its measured effect on protected queries, and the checks that authorize the database update. The prototype rejects changes that fail those checks. It has not yet shown that a partial upgrade can match a full upgrade while using less total encoding work.')
h('5. Objective(s) of Invention')
for x in ['Find search regressions before a partial model upgrade becomes active.','Keep enough information to identify the source and experiment state used in each preview.','Reject a commit when the preview is stale, altered or fails the current checks.','Count the target vectors computed during discovery and testing, including work on rejected actions.','Allow a checked legacy route and restoration of the latest unchanged committed membership.','Investigate whether this control process can reduce upgrade work without losing the benefit of the new model. This remains a research objective.']:p('• '+x)
h('6. Working principle of the invent (in brief)')
p('Let A be the old embedding model and B the new model. The controller computes B vectors for a proposed set of documents. Without changing live membership, it calculates the search results that would follow if those documents used B and the rest used A. It combines ranks from the two spaces and checks whether each protected result stays within its allowed rank. An optional development-query check also requires nDCG@10 to stay above a configured minimum.')
p('A SHA-256 digest identifies the preview record. At commit, the controller starts a SQLite write transaction, checks the source snapshot and record digest, and repeats the retrieval checks. It then writes the target vectors, active membership, epoch and commit history together. If a check or write fails, the membership change is not committed. The digest detects changes to that record; it is not a digital signature or a tamper-proof audit system.')
p('The prototype offers two ways to acquire candidates: a heuristic that examines a subset, and an exhaustive mode that computes every current B document vector. Strict admission refuses the heuristic evidence. Exhaustive mode establishes complete vector availability, but pays the full encoding cost and still does not guarantee that mixed results equal full-B results.')
page()
h('7. Description of the invention in detail (Include drawing and or photograph as needed)')
p('Figure 1 shows the implemented sequence. A preview asks what the search results would be after a particular change. Computing vectors and recording the answer do not activate the change.')
im=Image.new('RGB',(1400,540),'white');dr=ImageDraw.Draw(im)
try:font=ImageFont.truetype('/System/Library/Fonts/Supplemental/Arial.ttf',24)
except OSError:font=ImageFont.load_default()
boxes=[(30,40,430,155,'1  Source and policy snapshot\nIDs, hashes, epoch, models'),(500,40,900,155,'2  Candidate target vectors\nCache and charge all B work'),(970,40,1370,155,'3  Mixed serving preview\nNative ranks and action effects'),(970,330,1370,470,'4  Barrier and freshness check\nReject on failed checks'),(500,330,900,470,'5  Atomic durable commit\nVectors, membership, ledger'),(30,330,430,470,'6  Serve and retain history\nValidated fallback or rollback')]
for x1,y1,x2,y2,txt in boxes:dr.rounded_rectangle((x1,y1,x2,y2),12,fill='#f0f3f4',outline='black',width=2);dr.multiline_text((x1+15,y1+28),txt,font=font,fill='black',spacing=10)
for a,b in [((430,95),(500,95)),((900,95),(970,95)),((1170,155),(1170,330)),((970,400),(900,400)),((500,400),(430,400))]:
 dr.line([a,b],fill='black',width=4);x,y=b
 if b[0]>a[0]: pts=[(x,y),(x-12,y-7),(x-12,y+7)]
 elif b[0]<a[0]:pts=[(x,y),(x+12,y-7),(x+12,y+7)]
 else:pts=[(x,y),(x-7,y-12),(x+7,y-12)]
 dr.polygon(pts,fill='black')
im.save(OUT/'architecture.png');d.add_picture(str(OUT/'architecture.png'),width=Inches(6.4));d.add_paragraph('Figure 1  Implemented migration validation and commit sequence','Caption')
d.add_heading('Data and serving state',2)
p('Each document stores its ID, text, version, source hash, A vector, optional B vector and active membership. The B cache uses the ID and source hash. Reopening checks model identifiers and fusion policy, assuming unchanged model files. Preview and event records retain the evidence.')
p('In the default embodiment, migrated objects are excluded from the A result list. Reciprocal rank fusion uses 1/(60 + rank); ties favor B and then document ID. No cosine comparison is made between A and B vector spaces. The default fixture deliberately uses different vector dimensions, 128 and 256.')
d.add_heading('Preview and admission flow',2)
p('Figure 2 separates the proposed state from the active state. Candidate analysis may compute and cache B vectors, but it leaves active membership unchanged. The preview evaluates exactly the proposed member set and records the ranking changes for every declared protected query. A failed check ends that proposed action. A passing check creates an evidence record; it does not by itself activate the change.')
flow=Image.new('RGB',(1400,660),'white');fd=ImageDraw.Draw(flow)
steps=[('A  Current state','Documents, A/B membership, epoch'),('B  Propose action','Select document IDs; compute B vectors'),('C  Simulate retrieval','Rank A and B separately; fuse ranks'),('D  Test barrier','Required ranks and optional development gate'),('E  Bind preview','IDs, vectors, source hashes, policy, digest'),('F  Commit gate','Check freshness and constraints in write transaction')]
for i,(title,detail) in enumerate(steps):
 x=40+((i%3) if i<3 else (5-i))*460;y=30+(i//3)*310
 fd.rounded_rectangle((x,y,x+410,y+170),14,fill='#f4f4f4',outline='#222',width=3)
 fd.text((x+14,y+18),title,font=font,fill='black')
 fd.multiline_text((x+14,y+72),textwrap.fill(detail,26),font=font,fill='black',spacing=5)
for ax,ay,bx,by in [(450,115,500,115),(910,115,960,115),(1165,200,1165,340),(960,425,910,425),(500,425,450,425)]:
 fd.line((ax,ay,bx,by),fill='#222',width=4)
 if by>ay:fd.polygon([(bx,by),(bx-9,by-13),(bx+9,by-13)],fill='#222')
 elif bx>ax:fd.polygon([(bx,by),(bx-13,by-9),(bx-13,by+9)],fill='#222')
 else:fd.polygon([(bx,by),(bx+13,by-9),(bx+13,by+9)],fill='#222')
flow_path=OUT/'preview_commit_flow.png';flow.save(flow_path)
d.add_picture(str(flow_path),width=Inches(6.5));d.add_paragraph('Figure 2  Proposed action, preview evidence and commit gate. A rejected action leaves active membership unchanged.','Caption')
d.add_heading('Stored record and decision conditions',2)
table(['Item','Meaning in the prototype','Why it matters'],[
('Source version and hash','Version of each source document used to build its vector.','An edited document invalidates its earlier B representation.'),
('Active membership and epoch','Which representation serves each document, with a migration counter.','A preview for an older epoch cannot silently authorize a later state.'),
('Preview action and digest','Candidate IDs, vectors, tested results and SHA-256 record identifier.','The commit step detects alteration of the saved preview record.'),
('Protected query rule','Required document ID and maximum permitted rank.','A passing action must keep the declared result visible within its bound.'),
('Development gate','Optional minimum nDCG@10 on declared development queries.','Adds a broader relevance check, but only for those queries.')])
d.add_heading('Witness selection and commit',2)
p('Candidate discovery combines protected-query results, token overlap and required documents. The controller tests singletons and pairs over at most 24 candidates, then a protected-document bundle if needed. Each action record contains IDs, rank changes, pass status and object count. It chooses the smallest passing action among those tested, not a global or compute-cost optimum.')
p('The snapshot covers source versions, membership, model identifiers, fusion and the development gate. Commit rechecks the snapshot, digest, source hashes and current retrieval constraints within a write transaction. Tests cover altered evidence, stale source or epoch, and write failure. Source changes invalidate B. Rollback restores the latest unchanged committed membership; the demonstration uses a server that serializes requests.')
d.add_heading('Transaction and refusal cases',2)
p('The commit begins with a SQLite BEGIN IMMEDIATE transaction. It reads the current database state, compares it with the preview snapshot, verifies the evidence identifier and the current source hashes, and evaluates the retrieval barrier again. Only then does it update vector cache entries, active membership, the epoch and the event record. If any condition fails, the transaction rolls back. This ordering makes a preview an admission request rather than authority to mutate the index.')
table(['Observed case','Decision in the local prototype'],[
('Selected document source changes after preview','Reject stale evidence; invalidate B representation.'),
('Migration epoch changes after preview','Reject the earlier preview.'),
('Preview evidence is altered','Reject digest mismatch.'),
('A protected result moves below its bound','Reject that candidate action.'),
('A tested write fails','Roll back the local transaction; keep active membership unchanged.')])
d.add_heading('Example of an indirect ranking change',2)
p('In the authored fixture, moving car-guide to B leaves q-car within its allowed rank but moves password-policy to rank 2 for q-policy, whose maximum is 1. The controller therefore rejects the action. Moving password-policy instead passes all four checks. This illustrates why the preview examines the mixed result list, including effects on documents outside the selected action.')
d.add_heading('Separate experimental embodiment',2)
p('Continuity fusion retains all A vectors as a shadow ranking and weights A by 1 minus the migrated fraction and B by the migrated fraction. Its endpoints are legacy A and full B. This departs from the frozen A-removal embodiment and adds legacy storage and query work. ANN graph repair, distributed transactions and sublinear certified candidate discovery are not implemented.')
from disclosure_detail import add_detail, add_validation
add_detail(d,p,table)
page()
h('8. Experimental validation results:')
p('The prototype passed 25 automated regression tests. The browser demonstration also exercised candidate analysis, rejection, commit and rollback. These checks support the behaviour described below. They do not establish novelty, production reliability or a reduction in migration cost.',True)
table(['Experiment','Measured result','Interpretation'],[
('Authored 40-document fixture\n4 protected queries','14 B document encodings; 1 migrated object; 13 rejected singleton actions; 0 protected violations.','The gate makes an observable decision. Held-out full-B top-5 overlap is 0.80, below legacy 0.95.'),
('Full fixture trajectory','40 safe commits; 40 B document encodings; full membership migration.','The controller can finish this fixture. There is no final encoding saving.'),
('Exhaustive strict evidence\n40-document fixture','All 40 B vectors materialized and charged; an admissible commit completed.','Complete candidate coverage is available at full cost. Candidate enumeration remains capped.'),
('Source and transaction tests','Stale source, stale epoch, altered preview and injected write failure rejected.','Active membership is protected under the tested local SQLite failure cases.')])
d.add_heading('Public dataset experiments',2)
p('SciFact samples contain 160 documents, selected by retaining relevant documents for sampled queries and adding distractors. This relevance-aware selection can make evaluation easier and is not the full BEIR benchmark. The table reports held-out nDCG@10; higher values indicate better relevance ranking. L6 and L12 refer to the cached all-MiniLM sentence-transformer models.')
table(['Variant and split','A / selective / B','B charged / migrated','Result'],[
('L6 to L12, RRF\n6 protected, 12 held-out','.887 / .526 / .893','80 / 6','300 candidate previews rejected; 0 protected violations; 15% full-B top-10 envelope miss rate.'),
('L6 to L12, continuity\n6 protected, 12 held-out','.887 / .887 / .893','80 / 1','Preserved A quality without measured improvement. Exploratory split reused during redesign.'),
('Lexical to L12, continuity\n6 protected, 6 development, 6 held-out','.384 / .384 / .898','83 / 1','Development floor passed; target-quality progress was not achieved. Envelope miss rate 28.3%.')])
p('Full-B evaluation additionally encoded all 160 documents in a separate ledger for each run. It was not supplied to action selection. Historical runs differ in constraints and splits; they are not a controlled ablation. The random fixture baseline matches migrated-object count, not discovery compute. No equal-budget superiority or statistically significant quality gain is claimed.')
p('The experiments support a limited conclusion: the controller can reject actions that fail its declared checks. Passing those checks did not reliably preserve quality on other queries. The intended quality and cost advantage remains unproven.',True)
add_validation(d,p,table)
d.add_heading('Fresh reproducible prototype trace',2)
proof=json.loads((ROOT/'artifacts_private/prototype-proof-trace.json').read_text())
seq=proof['sequence'];committed=seq['commit']['state'];rolled=seq['rollback']
p('A separate script reran the functional sequence from a fresh SQLite database. It performs assertions rather than copying a dashboard reading. The raw trace is packaged as prototype-proof-trace.json; the 25-test transcript is packaged as test-results.txt.')
table(['Checkpoint','Measured state or result','What it demonstrates'],[
('Initial state','Epoch 0; active B 0; B document encodings 0.','No migration was active before the exercise.'),
('Candidate analysis','Epoch 0; active B 0; '+str(seq['before_commit']['charged_b_document_encodings'])+' B document encodings charged.','Discovery has real cost while membership stays unchanged.'),
('Unsafe action','car-guide; barrier fails; commit raises '+seq['unsafe']['commit_error']+'.','An inadmissible preview cannot activate membership.'),
('Passing action','password-policy; preview top five exactly equals the top five served after commit.','Preview and committed serving agree for the checked query in this run.'),
('Commit and restart','Epoch '+str(committed['epoch'])+'; active B '+str(committed['active_b'])+'; persisted snapshot agrees after reopening.','The accepted state is durable in local SQLite.'),
('Rollback','Epoch '+str(rolled['epoch'])+'; active B '+str(rolled['active_b'])+'; B encodings still charged '+str(rolled['charged_b_document_encodings'])+'.','Membership is restored without erasing incurred compute.'),
('Strict and fault modes','Heuristic strict path rejects; exhaustive strict path charges 40 B documents and commits; injected write failure rolls back.','The evidence policy and failure path affect control decisions.')])
p('The trace supports repeatable behaviour of this exact local fixture. It does not prove that every query is protected, that a real ANN index would behave the same way, or that the migration saves total compute.')
im=Image.new('RGB',(1400,370),'white');dr=ImageDraw.Draw(im)
stages=[('START','epoch 0  |  B active 0'),('ANALYSE','14 B encodings  |  B active 0'),('COMMIT','epoch 1  |  B active 1'),('ROLL BACK','epoch 2  |  B active 0')]
for i,(title,detail) in enumerate(stages):
 x=20+i*350;dr.rounded_rectangle((x,65,x+320,250),12,fill='#f4f4f4',outline='black',width=3)
 dr.text((x+18,95),title,font=font,fill='black');dr.multiline_text((x+18,157),textwrap.fill(detail,18),font=font,fill='black',spacing=4)
 if i<3:
  dr.line((x+320,157,x+350,157),fill='black',width=4);dr.polygon([(x+350,157),(x+337,149),(x+337,165)],fill='black')
proof_fig=OUT/'prototype_proof_sequence.png';im.save(proof_fig)
d.add_picture(str(proof_fig),width=Inches(6.5))
d.add_paragraph('Figure 3  Measured membership and charged target encodings during one fresh local proof run. The rejected action occurs before the commit checkpoint.','Caption')
d.add_heading('Reproducibility and evidence record',2)
p('The saved evidence package includes the 25-test log, a fixture benchmark, a 40-step migration trajectory, strict exhaustive evidence, three SciFact experiment records and a browser evidence export. The screenshots below are from the authored 40-document fixture. Their source JSON and the full-size captures are included in the package. The SciFact tables represent separate sampled experiments, not the pictured browser state.')
p('To reproduce the local browser case, start the server with a fresh database, analyze candidates, inspect the rejected car-guide action, commit the passing password-policy action, export the record and then roll back. The run instructions and exact command are in the accompanying demo guide. Repeating the historical SciFact measurements requires the recorded model snapshots and sampled IDs.')
page()
d.add_heading('Prototype screenshots',2)
p('Actual local browser captures from the authored lexical fixture. Screenshots illustrate the workflow; numeric results and failure checks are supported by the JSON exports and test log.')
for fname,caption in [('final-overview.jpg','Figure 4  Successful commit: 1 active B object, 14 target encodings and epoch 1.'),('final-blocked-detail.jpg','Figure 5  After rollback to epoch 2, the car-guide preview fails q-policy and commit is disabled.')]:
 image=Image.open(ROOT/'artifacts_private'/fname);w,hg=image.size
 width=min(6.7,7.0*w/hg)
 d.add_picture(str(ROOT/'artifacts_private'/fname),width=Inches(width))
 d.add_paragraph(caption,'Caption')
 if fname=='final-overview.jpg':
  page()
p('The full-resolution screenshots are included in the evidence package. No screenshot represents a neural benchmark or a production deployment.')
page()
d.add_page_break()
h('9. What aspect(s) of the invention need(s) protection?')
p('We ask the IPR Cell to examine the following sequence as a possible combined contribution. The proposed scope is limited to the concrete migration process described here:',True)
for x in ['Calculate the mixed search state for a selected set of B vectors, including removal of those documents from the active A ranking, before committing the change.','Record the effect of that action on protected result ranks, together with document IDs, source digests and experiment state, and use the record to choose an admissible action.','Identify the preview by a digest of its tested vectors and recorded state. Within the write transaction, reject stale or altered evidence and recheck the current retrieval constraints before changing membership.','In the strict mode, require current target-vector coverage before admission. The implemented exhaustive fallback computes all vectors and records that full cost.','Keep an inspectable event history, restore the latest unchanged committed membership when requested, and allow a legacy route only for a query that has passed its legacy check.']:p('• '+x)
p('The intended effect is to stop an index change that fails the inspected retrieval constraints or relies on stale source evidence. The local tests support that behaviour. Whether the combination adds an inventive step over model validation and standard transaction control remains unresolved. Rank fusion, hashing, caching, SQLite atomicity and the dashboard are supporting techniques. The exhaustive fallback is not claimed to save encoding work.')
p('A patent professional should assess the combined mechanism under the applicable computer-related invention requirements, including Section 3(k) of the Indian Patents Act [R7]. These proposed protection points are disclosure material, not final patent claims or a conclusion of patentability.')
d.add_heading('Technical feature comparison for claim review',2)
table(['Feature to examine','Closest established teaching','TopoPatch evidence and remaining question'],[
('Validate before activating a newer model','Amazon R8 validates model migration using client test data.','The local controller validates document-level mixed membership. Is that interaction technically distinct enough?'),
('Selectively reindex documents','Cloudera R2 updates affected vector-index content; HubSpot R1 supports model-related re-embedding.','The action is chosen from observed mixed-rank effects. Search scope is capped and savings are unproven.'),
('Keep old and new spaces available','Qdrant R3 documents practical embedding-model migration.','The tested path fuses native ranks and masks a selected A membership during preview.'),
('Check source and commit atomically','Checksums and ordinary database transactions are established.','The local prototype links a previewed retrieval result to the state checked at commit; inventive step remains open.'),
('Repair ANN topology','SPFresh R10 already studies local index updates.','No ANN graph repair is implemented here; it is outside the demonstrated protection request.')])
p('The possible point of distinction is the cooperation among a concrete mixed-state action preview, a query-level refusal decision and freshness-checked commit. Each ingredient alone appears in related work. The table is a claim-review aid, not a conclusion that the combination is new or non-obvious.')
d.add_page_break()
d.add_heading('References for sections 3 and 8',2)
refs=[
('R1','HubSpot. Vector as a service. US20240303516A1.','https://patents.google.com/patent/US20240303516A1/en'),
('R2','Cloudera. Methods and apparatus to manage input data sets to reflect dataset mutations for GenAI and RAG applications. US20250342179A1.','https://patents.google.com/patent/US20250342179A1/en'),
('R3','Qdrant. Migrate to a New Embedding Model.','https://qdrant.tech/documentation/tutorials-operations/embedding-model-migration/'),
('R4','BEIR benchmark repository and SciFact archive.','https://github.com/beir-cellar/beir'),
('R5','VIT. Intellectual Property Rights and Technology Transfer Policy and Guidelines. Amendment 16 April 2024.','https://vit.ac.in/files/IPR-and-Technology-Transfer-Policy.pdf'),
('R6','VIT. A guide on Technology Readiness Levels.','https://vit.ac.in/wp-content/uploads/2025/08/A_guide_on_Technology_Readiness_Levels.pdf'),
('R7','IP India. Patents Act 1970, Section 3.','https://ipindia.gov.in/acts/patent-act-1970/section-3'),
('R8','Amazon. Client-specific validation for migrating to a version of a machine learning model using client-supplied test data. US11610143B1.','https://patents.google.com/patent/US11610143B1/en'),
('R9','Embedding Patches for Localized Improvement of Embeddings. US20260195572A1.','https://patents.justia.com/patent/20260195572'),
('R10','Xu et al. SPFresh. SOSP 2023. Author laboratory publication page.','https://ustc-mlsys.github.io/publications/spfresh-sosp-2023/'),
('R11','SQLite. Isolation in SQLite.','https://sqlite.org/isolation.html')]
refs.append(('R12','Cormack, Clarke and Buettcher. Reciprocal Rank Fusion outperforms Condorcet and individual Rank Learning Methods. SIGIR 2009.','https://research.google/pubs/reciprocal-rank-fusion-outperforms-condorcet-and-individual-rank-learning-methods/'))
for n,t,url in refs:
 pp=p(f'[{n}] {t}\n{url}');pp.paragraph_format.space_after=Pt(5)
 for r in pp.runs:r.font.size=Pt(8)
p('Online references checked 15 October 2026. Source dataset archive: data/scifact.zip. The original project specification is retained in doc/.')
d.add_page_break()
blank=PdfReader(ROOT/'doc/Invention_Disclosure_Format_B.pdf').pages[1]
buf=BytesIO();mark=canvas.Canvas(buf,pagesize=(float(blank.mediabox.width),float(blank.mediabox.height)))
mark.setFont('Helvetica-Bold',18);mark.drawCentredString(185,446,'X');mark.save();buf.seek(0)
blank.merge_page(PdfReader(buf).pages[0])
marked_pdf=ROOT/'artifacts_private/vit-trl3-marked.pdf';w=PdfWriter();w.add_page(blank)
with marked_pdf.open('wb') as stream:w.write(stream)
pdftoppm='/Users/shaurya/.cache/codex-runtimes/codex-primary-runtime/dependencies/bin/override/pdftoppm'
subprocess.run([pdftoppm,'-f','1','-l','1','-scale-to','1800','-png',str(marked_pdf),str(ROOT/'artifacts_private/vit-trl3-marked')],check=True,stderr=subprocess.DEVNULL)
trl_image=ROOT/'artifacts_private/vit-trl3-marked-1.png'
cropped=ROOT/'artifacts_private/vit-trl3-section.png'
source_image=Image.open(trl_image)
source_image.crop((0,0,source_image.width,int(source_image.height*.57))).save(cropped)
d.add_picture(str(cropped),width=Inches(6.85))
d.core_properties.title='TopoPatch invention disclosure';d.core_properties.subject='VIT Invention Disclosure Format B';d.core_properties.author=''
d.save(OUT/'TopoPatch_IDF_B.docx')
